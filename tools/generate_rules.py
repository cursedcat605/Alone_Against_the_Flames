"""从《向火独行 卡.xlsx》生成引擎需要的规则数据表。

生成：
  src/content/rules/occupations.json   职业模板（115 个）
  src/content/rules/specialties.json   技能专攻分支的基础值

不生成（人工整理的小表，直接手写在 src/content/rules/ 里）：
  skills.json   基础技能值与基础成功率
  rules.json    属性公式、衍生属性、年龄调整、伤害加值表、快速参考规则

设计取舍：只在 xlsx 中结构**规整**时才自动提取（职业列表/技能专攻都是干净的二维表），
其余手写。这样规则数据可审计、可重跑，又不至于为了几个数字写脆弱的解析器。

跑法：python tools/generate_rules.py
"""

from __future__ import annotations

import json
import re
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "向火独行 卡.xlsx"
OUT_DIR = ROOT / "src" / "content" / "rules"


def cell_text(value) -> str:
    """单元格转文本；整数化的 float 去掉小数点（xlsx 里 30.0 要变 '30'）。"""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def parse_credit(raw: str) -> dict:
    m = re.match(r"^\s*(\d+)\s*[-–—~]\s*(\d+)\s*$", raw)
    if not m:
        return {"min": 0, "max": 99, "raw": raw}
    return {"min": int(m.group(1)), "max": int(m.group(2))}


def parse_int(raw: str, default: int = 0) -> int:
    m = re.search(r"\d+", raw)
    return int(m.group()) if m else default


def split_top_level(text: str) -> list[str]:
    """只在括号**之外**按分隔符切分。

    本职技能里大量出现 "两项社交技能（魅惑，话术，恐吓，说服）" 这种写法。
    无脑 re.split 会把括号里的候选项也切出来，于是「话术」「恐吓」「说服）」
    各自变成一条独立的「本职技能」，静默拿到职业点。
    """
    depth = 0
    buf: list[str] = []
    out: list[str] = []
    for ch in text:
        if ch in "（(":
            depth += 1
        elif ch in "）)":
            depth = max(0, depth - 1)
        if depth == 0 and ch in "，,、；;":
            out.append("".join(buf))
            buf = []
            continue
        buf.append(ch)
    out.append("".join(buf))
    return [
        s.strip().rstrip("。.；;")
        for s in out
        if s.strip().rstrip("。.；;")
    ]


def gen_occupations(wb) -> dict:
    ws = wb["职业列表"]
    items = []
    for r in range(2, ws.max_row + 1):
        num = ws.cell(r, 1).value
        if num is None:
            continue
        name = cell_text(ws.cell(r, 2).value)
        if not name:
            continue
        raw_credit = cell_text(ws.cell(r, 3).value)
        raw_skills = cell_text(ws.cell(r, 6).value)
        items.append({
            "id": int(num),
            "name": name,
            "creditRating": parse_credit(raw_credit),
            "attrFormula": cell_text(ws.cell(r, 4).value),
            "skillPoints": parse_int(cell_text(ws.cell(r, 5).value)),
            # 本职技能是自然语言（"任意两项其他个人或时代特长"），引擎只做展示与提示，
            # 不强制约束玩家分配，因此原样保留字符串数组（但要按顶层分隔符切）。
            "occupationSkills": split_top_level(raw_skills),
        })
    return {
        "meta": {
            "source": "向火独行 卡.xlsx / 职业列表",
            "count": len(items),
            "note": "skillPoints 为职业（本职）技能点总数；兴趣技能点 = 智力×2，见 rules.json",
        },
        "occupations": items,
    }


SPECIALTY_GROUPS = [
    ("艺术与手艺", "arts", 2, 3),
    ("科学", "science", 5, 6),
    ("格斗", "fighting", 8, 9),
    ("射击", "firearms", 11, 12),
    ("非常规技能", "uncommon", 14, 15),
]


def gen_specialties(wb) -> dict:
    ws = wb["技能专攻"]
    groups = []
    duplicates: list[str] = []
    for label, key, col_name, col_base in SPECIALTY_GROUPS:
        items = []
        seen: set[str] = set()
        for r in range(4, ws.max_row + 1):
            name = cell_text(ws.cell(r, col_name).value)
            raw_base = cell_text(ws.cell(r, col_base).value)
            if not name or not raw_base.isdigit():
                continue
            if name in seen:
                # 表里「表演」出现了两次。重复项会让建卡界面列出两个同名专攻，
                # 玩家选中哪一个都看不出区别。去重，但要说出来 —— 静默丢弃
                # 和被丢弃的东西本身一样可疑。
                duplicates.append(f"{label}/{name}")
                continue
            seen.add(name)
            items.append({"name": name, "base": int(raw_base)})
        groups.append({"group": label, "key": key, "items": items})
    if duplicates:
        print(f"  专攻表里有重复项，已保留首次出现：{'、'.join(duplicates)}")
    return {
        "meta": {"source": "向火独行 卡.xlsx / 技能专攻"},
        "groups": groups,
    }


def write(path: Path, payload: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"OK 写入 {path.relative_to(ROOT)}")


def main() -> None:
    if not XLSX.exists():
        raise SystemExit(f"找不到规则表：{XLSX}")
    wb = load_workbook(XLSX, data_only=True)
    occ = gen_occupations(wb)
    spe = gen_specialties(wb)
    write(OUT_DIR / "occupations.json", occ)
    write(OUT_DIR / "specialties.json", spe)
    print(f"职业 {occ['meta']['count']} 个；技能专攻 {sum(len(g['items']) for g in spe['groups'])} 项")


if __name__ == "__main__":
    main()
