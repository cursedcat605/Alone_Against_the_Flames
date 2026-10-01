"""一次性用的：把「向火独行 卡.xlsx」的所有工作表内容原样 dump 出来。

用途：摸清建卡规则表的结构，据此设计 rules/*.json。
跑法：python tools/inspect_xlsx.py
"""

from __future__ import annotations

import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:  # pragma: no cover
    print("MISSING_OPENPYXL")
    sys.exit(2)

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "向火独行 卡.xlsx"


def cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def main() -> None:
    # 注意：不要用 PowerShell 的 `>` 重定向，PS 5.1 会写成 UTF-16，
    # 导致文件被当成二进制。直接在 Python 里按 UTF-8 落盘。
    out_path = ROOT / "tools" / "_xlsx_dump.txt"
    lines: list[str] = []
    lines.append(f"FILE={XLSX}")
    lines.append(f"EXISTS={XLSX.exists()}")
    if not XLSX.exists():
        out_path.write_text("\n".join(lines), encoding="utf-8")
        sys.exit(1)

    wb = openpyxl.load_workbook(XLSX, data_only=True)
    lines.append(f"SHEETS={wb.sheetnames}")
    for ws in wb.worksheets:
        lines.append("")
        lines.append("=" * 70)
        lines.append(f"SHEET: {ws.title}  dims={ws.dimensions}  "
                     f"max_row={ws.max_row} max_col={ws.max_column}")
        lines.append("=" * 70)
        for row in ws.iter_rows():
            cells = [cell_text(c.value) for c in row]
            if not any(cells):
                continue
            # 保留列位置信息，方便还原表格结构
            parts = [
                f"[{c.coordinate}]{cells[i]}"
                for i, c in enumerate(row)
                if cells[i]
            ]
            lines.append(" | ".join(parts))

    out_path.write_text("\n".join(lines), encoding="utf-8")
    print(f"WROTE {out_path} ({len(lines)} lines)")
    print("SHEETS=" + ", ".join(wb.sheetnames))


if __name__ == "__main__":
    main()
